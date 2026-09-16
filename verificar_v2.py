# -*- coding: utf-8 -*-
import docx
doc = docx.Document(r'C:\Users\mabad\Desktop\Avance de Tesis (1)_CORREGIDA.docx')
print(f'Total parrafos: {len(doc.paragraphs)}')
texts = [p.text for p in doc.paragraphs]
checks = [
    'Ingeniería Civil Informática',
    'Pregunta de Investigación',
    'Hipótesis',
    'Python',
    'Ultralytics YOLOv8',
    'OpenCV',
    'Desarrollar e implementar un sistema',
    'data augmentation',
    'hiperparámetros',
    'garantizando baja latencia',
]
for c in checks:
    found = any(c in t for t in texts)
    print(f'  {c[:40]}: {"OK" if found else "FALTA"}')
