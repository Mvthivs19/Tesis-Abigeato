# -*- coding: utf-8 -*-
"""Add missing tools to the thesis document."""
import docx
from docx.oxml.ns import qn

INPUT = r"C:\Users\mabad\Desktop\Avance de Tesis (1)_EDITADA.docx"
OUTPUT = r"C:\Users\mabad\Desktop\Avance de Tesis (1)_EDITADA.docx"

doc = docx.Document(INPUT)

# Find the paragraph with "Ultralytics YOLOv8" (last tool added) to insert after it
yolo_idx = None
for i, p in enumerate(doc.paragraphs):
    if "Ultralytics YOLOv8" in p.text and "implementación oficial" in p.text:
        yolo_idx = i
        break

if yolo_idx is None:
    print("ERROR: No se encontró el párrafo de Ultralytics YOLOv8")
    exit(1)

print(f"Encontrado Ultralytics YOLOv8 en párrafo {yolo_idx}")

# Tools to add after YOLOv8
tools = [
    ("OpenCV", "OpenCV (Open Source Computer Vision Library) es la biblioteca estándar para el procesamiento de imágenes y video en tiempo real. Proporciona funciones optimizadas para operaciones morfológicas, manipulación de color, detección de bordes y manejo de flujos de video. En este proyecto, OpenCV cumple tres roles fundamentales: la captura y decodificación del feed de video, la simulación del espectro NIR mediante transformaciones de escala de grises y corrección gamma, y la renderización de la interfaz visual con superposición de bounding boxes y polígonos de vigilancia (Bradski, 2000)."),
    ("FiftyOne", "FiftyOne es una herramienta de código abierto diseñada para la construcción, análisis y gestión de datasets de visión por computadora. Permite descargar datasets públicos como Open Images V7, inspeccionar distribuciones de clases y filtrar anotaciones. En este proyecto, FiftyOne se emplea para la descarga selectiva de las imágenes de las cinco clases objetivo del dataset Open Images V7, facilitando la etapa de recolección y preprocesamiento del material de entrenamiento (Voxel51, 2024)."),
    ("SQLite", "SQLite es un sistema de gestión de bases de datos relacional embebido, es decir, no requiere un servidor separado para funcionar. Almacena toda la base de datos en un único archivo binario, lo que lo hace ideal para dispositivos Edge con recursos limitados. En este proyecto, SQLite almacena la tabla de eventos de detección, el log de salud del sistema y la cola de notificaciones remotas pendientes, garantizando persistencia local sin infraestructura adicional (Hipp et al., 2024)."),
    ("Scikit-learn", "Scikit-learn es una biblioteca de Python para machine learning que proporciona herramientas para clasificación, regresión, clustering y preprocesamiento de datos. En este proyecto se utiliza específicamente la función train_test_split para la división estratificada del dataset en los conjuntos de entrenamiento, validación y prueba (80/10/10), asegurando una distribución equilibrada de clases en cada partición (Pedregosa et al., 2011)."),
    ("Draw.io", "Draw.io es una herramienta gratuita y de código abierto para la creación de diagramas, incluyendo diagramas de flujo, UML, organigramas y representaciones de arquitectura de software. Su interfaz basada en navegador permite exportar diagramas en múltiples formatos. En este proyecto se utiliza para el diseño de los diagramas de la arquitectura de vistas 4+1, los diagramas de componentes y el diagrama de despliegue del sistema (KeepCoding, 2025)."),
    ("Bizagi Modeler", "Bizagi Modeler es una herramienta de modelado de procesos de negocios que utiliza la notación estándar BPMN (Business Process Model and Notation). Permite diseñar, simular y documentar flujos de trabajo operativos. En este proyecto se emplea para diagramar el proceso completo de detección de intrusos, desde la captura del frame hasta la detonación de la alerta y el registro en base de datos, lo que facilita la auditoría del protocolo de seguridad (Bizagi, 2025)."),
]

ref_element = doc.paragraphs[yolo_idx]._element

for titulo, desc in tools:
    # Create new paragraph
    new_p = docx.oxml.OxmlElement('w:p')
    
    # Title in bold
    new_r_title = docx.oxml.OxmlElement('w:r')
    new_rPr = docx.oxml.OxmlElement('w:rPr')
    new_b = docx.oxml.OxmlElement('w:b')
    new_rPr.append(new_b)
    new_r_title.append(new_rPr)
    new_t_title = docx.oxml.OxmlElement('w:t')
    new_t_title.text = titulo
    new_r_title.append(new_t_title)
    new_p.append(new_r_title)
    
    # Description
    new_r_desc = docx.oxml.OxmlElement('w:r')
    new_t_desc = docx.oxml.OxmlElement('w:t')
    new_t_desc.text = '\n' + desc
    new_r_desc.append(new_t_desc)
    new_p.append(new_r_desc)
    
    ref_element.addnext(new_p)
    ref_element = new_p
    print(f"  Agregada: {titulo}")

doc.save(OUTPUT)
print(f"\nArchivo guardado: {OUTPUT}")
