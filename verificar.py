import docx
doc = docx.Document(r'C:\Users\mabad\Desktop\Avance de Tesis (1)_EDITADA.docx')

# Check all tools are present
tools = ['Python', 'PyTorch', 'Ultralytics YOLOv8', 'OpenCV', 'FiftyOne', 'SQLite', 'Scikit-learn', 'Draw.io', 'Bizagi Modeler']
for t in tools:
    found = any(t in p.text for p in doc.paragraphs)
    print(f'  {t}: {"OK" if found else "FALTA"}')

print()
checks = [
    'Ingeniería Civil Informática',
    'Desarrollar e implementar un sistema',
    'data augmentation',
    'hiperparámetros',
    'garantizando baja latencia',
    'encargado de ejecutar',
    'así identificar',
    'se tendrá',
    'Pregunta de Investigación',
    'Hipótesis',
]
for text in checks:
    found = any(text in p.text for p in doc.paragraphs)
    print(f'  "{text[:50]}": {"OK" if found else "FALTA"}')
