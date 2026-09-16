# -*- coding: utf-8 -*-
"""
Edición CORRECTA de la tesis - solo lo esencial, sin destruir estructura.
"""
import docx
from docx.oxml.ns import qn

INPUT = r"C:\Users\mabad\Desktop\Avance_Tesis_LIMPIA.docx"
OUTPUT = r"C:\Users\mabad\Desktop\Avance de Tesis (1)_CORREGIDA.docx"

doc = docx.Document(INPUT)

print("=== Corrigiendo erratas ===")
correcciones = {
    "Ingenieria Civil Informatica": "Ingeniería Civil Informática",
    "Ingenieria civil informatica": "Ingeniería civil informática",
    "Desarrollar e implemente un sistema": "Desarrollar e implementar un sistema",
    "en tiempo en tiempo real": "en tiempo real",
    "Data argumentation": "data augmentation",
    "hiperparametros": "hiperparámetros",
    "garantizado baja latencia": "garantizando baja latencia",
    "Es el grupo multidisciplinario de ejecutar": "Es el grupo multidisciplinario encargado de ejecutar",
    "permitiendo asi el identificar": "permitiendo así identificar",
    "que se tendra": "que se tendrá",
}

for p in doc.paragraphs:
    for old, new in correcciones.items():
        if old in p.text:
            for run in p.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    print(f"  OK: '{old[:35]}...' -> '{new[:35]}...'")

# ============================================================
# AGREGAR PREGUNTA DE INVESTIGACIÓN E HIPÓTESIS
# (Solo texto, sin tablas, antes de Objetivos)
# ============================================================
print("\n=== Agregando Pregunta de Investigación e Hipótesis ===")

idx_objetivos = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Objetivos" and p.style.name == "Heading 2":
        idx_objetivos = i
        break

if idx_objetivos:
    ref = doc.paragraphs[idx_objetivos]._element
    
    # Pregunta de Investigación - Heading 2
    h_preg = docx.oxml.OxmlElement('w:p')
    h_preg_pPr = docx.oxml.OxmlElement('w:pPr')
    h_preg_style = docx.oxml.OxmlElement('w:pStyle')
    h_preg_style.set(qn('w:val'), 'Heading2')
    h_preg_pPr.append(h_preg_style)
    h_preg.append(h_preg_pPr)
    h_preg_r = docx.oxml.OxmlElement('w:r')
    h_preg_t = docx.oxml.OxmlElement('w:t')
    h_preg_t.text = "Pregunta de Investigación"
    h_preg_r.append(h_preg_t)
    h_preg.append(h_preg_r)
    ref.addprevious(h_preg)
    
    # Párrafo con la pregunta
    p_preg = docx.oxml.OxmlElement('w:p')
    p_preg_r = docx.oxml.OxmlElement('w:r')
    p_preg_t = docx.oxml.OxmlElement('w:t')
    p_preg_t.text = "¿Es posible desarrollar un sistema de alerta temprana basado en visión computacional y procesamiento local (Edge Computing) que clasifique morfológicamente intrusos humanos y especies ganaderas en condiciones de visibilidad nocturna, alcanzando métricas de precisión superiores al 85% y una tasa de falsos positivos inferior al 5%?"
    p_preg_r.append(p_preg_t)
    p_preg.append(p_preg_r)
    h_preg.addnext(p_preg)
    
    # Hipótesis - Heading 2
    h_hip = docx.oxml.OxmlElement('w:p')
    h_hip_pPr = docx.oxml.OxmlElement('w:pPr')
    h_hip_style = docx.oxml.OxmlElement('w:pStyle')
    h_hip_style.set(qn('w:val'), 'Heading2')
    h_hip_pPr.append(h_hip_style)
    h_hip.append(h_hip_pPr)
    h_hip_r = docx.oxml.OxmlElement('w:r')
    h_hip_t = docx.oxml.OxmlElement('w:t')
    h_hip_t.text = "Hipótesis"
    h_hip_r.append(h_hip_t)
    h_hip.append(h_hip_r)
    p_preg.addnext(h_hip)
    
    # Párrafo con la hipótesis
    p_hip = docx.oxml.OxmlElement('w:p')
    p_hip_r = docx.oxml.OxmlElement('w:r')
    p_hip_t = docx.oxml.OxmlElement('w:t')
    p_hip_t.text = "Si se implementa un modelo de detección de objetos YOLOv8 optimizado mediante transfer learning, entrenado con imágenes NIR simuladas de las cinco clases objetivo, entonces el sistema será capaz de clasificar con precisión superior al 85% y una tasa de falsos positivos inferior al 5% la presencia de intrusos humanos frente a especies ganaderas, operando de forma autónoma en dispositivos Edge sin dependencia de conectividad a la nube."
    p_hip_r.append(p_hip_t)
    p_hip.append(p_hip_r)
    h_hip.addnext(p_hip)
    
    print("  OK: Pregunta de Investigación e Hipótesis agregadas")

# ============================================================
# REEMPLAZAR "Por definir" con herramientas reales
# ============================================================
print("\n=== Completando herramientas del Marco Teórico ===")

por_definir_indices = []
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "Por definir":
        por_definir_indices.append(i)

print(f"  Encontrados {len(por_definir_indices)} 'Por definir'")

herramientas = [
    ("Python", "Python es un lenguaje de programación interpretado, de alto nivel y propósito general, reconocido por su sintaxis clara y legible. Su amplio ecosistema de librerías lo convierten en el estándar de facto para proyectos de inteligencia artificial, ciencia de datos y automatización. En este proyecto, Python sirve como columna vertebral del sistema, integrando desde la preparación del dataset hasta la inferencia en tiempo real y el almacenamiento en base de datos (Python Software Foundation, 2024)."),
    ("Ultralytics YOLOv8", "Ultralytics YOLOv8 es la implementación oficial del algoritmo de detección de objetos en un solo paso YOLO en su octava versión. A diferencia de los modelos de dos etapas, YOLO localiza y clasifica los objetos en una única pasada de inferencia, lo que lo hace ideal para aplicaciones en tiempo real con restricciones de hardware. La variante YOLOv8n (nano) emplea apenas 3.2 millones de parámetros, permitiendo ejecutarse en dispositivos Edge de bajo consumo. Para este proyecto se utilizó transfer learning a partir de los pesos preentrenados en COCO (Jocher et al., 2023)."),
    ("OpenCV", "OpenCV es la biblioteca estándar para el procesamiento de imágenes y video en tiempo real. Proporciona funciones optimizadas para operaciones morfológicas, manipulación de color y manejo de flujos de video. En este proyecto cumple tres roles: captura y decodificación del feed de video, simulación del espectro NIR mediante transformaciones de escala de grises, y renderización de la interfaz visual con superposición de bounding boxes (Bradski, 2000)."),
]

for idx, i in enumerate(por_definir_indices):
    if idx >= len(herramientas):
        break
    titulo, desc = herramientas[idx]
    p = doc.paragraphs[i]
    for run in p.runs:
        run.text = ""
    run_title = p.add_run(titulo)
    run_title.bold = True
    run_desc = p.add_run("\n" + desc)
    print(f"  OK: {titulo}")

# Si hay menos "Por definir" que herramientas, agregar las restantes después del último
if len(por_definir_indices) < len(herramientas):
    last_idx = por_definir_indices[-1] if por_definir_indices else 218
    ref_elem = doc.paragraphs[last_idx]._element
    for titulo, desc in herramientas[len(por_definir_indices):]:
        new_p = docx.oxml.OxmlElement('w:p')
        new_r_title = docx.oxml.OxmlElement('w:r')
        new_rPr = docx.oxml.OxmlElement('w:rPr')
        new_b = docx.oxml.OxmlElement('w:b')
        new_rPr.append(new_b)
        new_r_title.append(new_rPr)
        new_t_title = docx.oxml.OxmlElement('w:t')
        new_t_title.text = titulo
        new_r_title.append(new_t_title)
        new_p.append(new_r_title)
        new_r_desc = docx.oxml.OxmlElement('w:r')
        new_t_desc = docx.oxml.OxmlElement('w:t')
        new_t_desc.text = '\n' + desc
        new_r_desc.append(new_t_desc)
        new_p.append(new_r_desc)
        ref_elem.addnext(new_p)
        ref_elem = new_p
        print(f"  OK: {titulo} (agregado)")

doc.save(OUTPUT)
print(f"\n=== Archivo guardado: {OUTPUT} ===")
